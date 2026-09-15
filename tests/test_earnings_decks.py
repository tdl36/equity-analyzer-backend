import copy
import io
import unittest
from earnings_decks import build, revise, render_pptx
from pptx import Presentation

ROW={'id':'bd77a316-24d7-41ac-a11c-c354fdbf873a','ticker':'DEMO','status':'completed','input':{'topic':'Illustrative quarterly review'},'output':{'synthesisMarkdown':'# Executive view\n- Revenue was 100 units.\n- Guidance is unchanged.\n- Analyst interpretation remains uncertain.\n## Open questions\nWhat evidence would resolve the margin debate?','sourceFiles':['Illustrative_release.pdf']}}

class EarningsDeckTests(unittest.TestCase):
    def test_full_retains_all_words_and_freezes_source(self):
        row=copy.deepcopy(ROW);row['output']['synthesisMarkdown']+='\n'+'long evidence '*6000+' FINAL_MARKER'
        deck=build(row)
        self.assertIn('FINAL_MARKER',' '.join(x for s in deck['slides'] for x in s['items']))
        self.assertEqual(deck['sources'][0]['filename'],'Illustrative_release.pdf')
        row['output']['synthesisMarkdown']='changed'
        self.assertNotEqual(deck['recapHash'],build(row)['recapHash'])
    def test_brief_discloses_omissions_retains_notes(self):
        deck=build(ROW,'brief');self.assertEqual(deck['omittedParagraphs'],1)
        self.assertIn('Analyst interpretation',' '.join(s['notes'] for s in deck['slides']))
        self.assertTrue(deck['warnings'])
    def test_no_recap_and_running_refused(self):
        for row in [{**ROW,'output':{}},{**ROW,'status':'running'}]:
            with self.assertRaises(ValueError):build(row)
    def test_source_edit_rejected_and_metadata_immutable(self):
        deck=build(ROW);edit=copy.deepcopy(deck);edit['sources']=[];edit['scope']='verified';edit['slides'][1]['items'][0]='My interpretation'
        saved=revise(deck,edit);self.assertEqual(saved['sources'],deck['sources']);self.assertEqual(saved['scope'],deck['scope']);self.assertTrue(saved['slides'][1]['edited'])
        edit['slides'][-1]['items']=['Invented original']
        with self.assertRaises(ValueError):revise(deck,edit)
    def test_oversized_edit_and_slide_identity_refused(self):
        deck=build(ROW);edit=copy.deepcopy(deck);edit['slides'][1]['items']=['x'*241]
        with self.assertRaises(ValueError):revise(deck,edit)
        edit=copy.deepcopy(deck);edit['slides'][0]['id']='unknown'
        with self.assertRaises(ValueError):revise(deck,edit)
    def test_long_original_filename_survives_save_and_export(self):
        row=copy.deepcopy(ROW);row['output']['sourceFiles']=['long-name-'*40+'.pdf']
        deck=build(row);self.assertEqual(revise(deck,deck)['sources'],deck['sources'])
        self.assertTrue(render_pptx(deck))
    def test_pptx_is_native_widescreen_and_carries_provenance(self):
        for theme in ['paper','midnight','sage']:
            deck=build(ROW,theme=theme);prs=Presentation(io.BytesIO(render_pptx(deck)))
            self.assertAlmostEqual(prs.slide_width/prs.slide_height,16/9,places=3)
            self.assertTrue(all(shape.shape_type!=13 for slide in prs.slides for shape in slide.shapes))
            self.assertIn(deck['recapHash'],prs.slides[1].notes_slide.notes_text_frame.text)
            self.assertIn('Revenue was 100',' '.join(shape.text for slide in prs.slides for shape in slide.shapes if shape.has_text_frame))
    def test_unknown_style_and_empty_edits_rejected(self):
        with self.assertRaises(ValueError):build(ROW,theme='unknown')
        with self.assertRaises(ValueError):revise(build(ROW),{})


# Exercise the actual Flask routes with a deliberately isolated database double.
from contextlib import contextmanager
from flask import Flask
from earnings_decks import create_blueprint
import json

class MemoryCursor:
    def __init__(self):self.rows=[];self.result=[]
    def execute(self,sql,args=()):
        if sql.startswith(('CREATE TABLE','SELECT pg_advisory')):self.result=[];return
        if 'FROM analyst_activities' in sql:self.result=[copy.deepcopy(ROW)] if args[0]==ROW['id'] else [];return
        if sql.startswith('INSERT INTO earnings_decks'):
            key=args[0];rev=1 if len(args)==3 else args[1];activity=args[1] if len(args)==3 else args[2];body=json.loads(args[-1])
            if not any(r['id']==key and r['revision']==rev for r in self.rows):self.rows.append({'id':key,'revision':rev,'activity_id':activity,'body':body,'created_at':'2026-09-15'})
            return
        if 'FROM earnings_decks' in sql:
            rows=[r for r in self.rows if r['activity_id']==args[0]] if 'WHERE activity_id=' in sql else [r for r in self.rows if r['id']==args[0]]
            if 'AND revision=' in sql:rows=[r for r in rows if r['revision']==args[1]]
            rows=sorted(rows,key=lambda r:r['revision'],reverse=True)
            if 'DISTINCT ON' in sql:rows=list({r['id']:r for r in reversed(rows)}.values())
            if 'LIMIT 1' in sql:rows=rows[:1]
            self.result=copy.deepcopy(rows);return
        raise AssertionError(sql)
    def fetchone(self):return self.result[0] if self.result else None
    def fetchall(self):return self.result

class DeckApiTests(unittest.TestCase):
    def setUp(self):
        self.cursor=MemoryCursor()
        @contextmanager
        def db(commit=False):yield None,self.cursor
        app=Flask(__name__);app.register_blueprint(create_blueprint(db));self.client=app.test_client()
        self.key='3e557fa9-1b56-439f-ae6d-04f536bd45f5'
    def create(self):return self.client.post('/api/earnings/decks',json={'activityId':ROW['id'],'requestId':self.key}).get_json()
    def test_idempotent_create_no_duplicate(self):
        first=self.create();second=self.create();self.assertEqual(first,second);self.assertEqual(len(self.cursor.rows),1)
        self.assertEqual(len(self.client.get('/api/earnings/decks?activityId='+ROW['id']).get_json()['decks']),1)
    def test_version_conflict_and_original_export(self):
        deck=self.create();body=copy.deepcopy(deck['body']);body['slides'][1]['items'][0]='Edited interpretation'
        response=self.client.put('/api/earnings/decks/'+self.key,json={'revision':1,'body':body});self.assertEqual(response.status_code,200)
        self.assertEqual(response.get_json()['revision'],2)
        self.assertEqual(self.client.put('/api/earnings/decks/'+self.key,json={'revision':1,'body':body}).status_code,409)
        old=self.client.get('/api/earnings/decks/'+self.key+'?revision=1').get_json();self.assertEqual(old['body'],deck['body']);self.assertEqual(len(old['history']),2)
        export=self.client.get('/api/earnings/decks/'+self.key+'/export?revision=1');self.assertEqual(export.status_code,200);self.assertTrue(export.data.startswith(b'PK'))
    def test_invalid_unknown_and_cross_event(self):
        self.assertEqual(self.client.get('/api/earnings/decks?activityId=no').status_code,400)
        self.assertEqual(self.client.get('/api/earnings/decks/'+self.key).status_code,404)
        self.create()
        response=self.client.post('/api/earnings/decks',json={'activityId':'edd717bd-2c02-4ed4-a03c-c3f0cd6c9b49','requestId':self.key})
        self.assertEqual(response.status_code,400)

class SavedEvidenceTests(unittest.TestCase):
    def test_saved_checks_and_comparisons_remain_explicit(self):
        row=copy.deepcopy(ROW)
        row['output']['claimReview']={'claims':[{'id':'1','passageMatched':True,'reviewPassed':True},{'id':'2','passageMatched':False,'reviewPassed':True}], 'numericComparisons':[{'metric':'Revenue','actual':{'value':100,'unit':'USD m','period':'Q2','basis':'GAAP'},'benchmark':{'value':99,'unit':'USD m','period':'Q2','basis':'GAAP'},'status':'arithmetic_checked','delta':1}]}
        d=build(row);text=' '.join(x for s in d['slides'] for x in s['items'])
        self.assertIn('1 of 2 selected claims passed',text)
        self.assertIn('100 · USD m · Q2 · GAAP',text)
        self.assertIn('Saved comparison status: arithmetic checked',text)
        self.assertEqual(ROW['output'].get('claimReview'),None)

if __name__=='__main__':unittest.main()
