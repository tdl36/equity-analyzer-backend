import copy, unittest, uuid
from pipeline_lab_note_provenance import validate, from_assignment, export_lines
from pipeline_lab_notes import normalized, edit, read
from tests.unit.test_pipeline_lab_notes import DraftVersions


def fixture():
    pack=[{'reference':'charlie:thesis_1','filename':'source.pdf','sha256':'a'*64,'bytes':10}]
    value={'version':1,'purpose':{'kind':'event_note','event':{'title':'Earnings <fixture>','date':'2025-03-01'}},'sourceEvidence':{'version':1,'sources':[{'id':'1','filename':'source.pdf','sha256':'a'*64,'role':'supplemental','publishedDate':'2024-03-01','publisher':'Fixture','sourceKind':'transcript','referencedBySavedThesis':None}]}}
    return pack,value


class Provenance(unittest.TestCase):
    def test_exact_source_validation_and_legacy(self):
        pack,value=fixture()
        self.assertEqual(validate(value,pack),value)
        self.assertEqual(from_assignment({},pack),{})
        bad=copy.deepcopy(value);bad['sourceEvidence']['sources'][0]['sha256']='b'*64
        with self.assertRaises(ValueError):validate(bad,pack)
        with self.assertRaises(ValueError):validate({**value,'version':True},pack)
        with self.assertRaises(ValueError):validate({**value,'purpose':{'kind':'inferred'}},pack)
        self.assertIsNone(validate({**value,'sourceEvidence':None},pack)['sourceEvidence'])

    def test_export_scope_and_safe_metadata(self):
        pack,value=fixture();value['purpose']['event']['title']='<script> [click](https://invalid.example)'
        text='\n'.join(export_lines(validate(value,pack)))
        self.assertIn('does not update the thesis',text);self.assertIn('Supplemental history',text)
        self.assertIn('2024',text);self.assertNotIn('<script>',text);self.assertNotIn('[click](',text)
        self.assertIn('not inferred', '\n'.join(export_lines(None)))
        value['purpose']={'kind':'thesis_review'}
        self.assertIn('Longitudinal thesis review','\n'.join(export_lines(value)))

    def test_word_export_retains_scope_and_plain_metadata(self):
        from pipeline_lab_note_review import export_text, word_export
        from docx import Document
        pack,value=fixture()
        draft={'title':'Fixture','markdown':'Test note','revision':1,'createdAt':'2026-10-15','sourcePack':pack,'qualityIssues':[],'researchProvenance':value}
        text=export_text(draft)
        doc=Document(word_export('ABT','Fixture',text))
        visible='\n'.join(p.text for p in doc.paragraphs)
        for item in ['Earnings <fixture>', '2025-03-01', '2024-03-01', 'Supplemental history', 'does not update the thesis']:
            self.assertIn(item,visible)
        for p in doc.paragraphs:
            for run in p.runs:
                self.assertEqual(run.font.name,'Calibri');self.assertEqual(str(run.font.color.rgb),'000000')


class VersionRetention(DraftVersions):
    def test_scope_survives_edits_and_cannot_be_substituted(self):
        self.payload['sourcePack'],self.payload['researchProvenance']=fixture()
        self.save();event=str(uuid.uuid4())
        with self.db(True) as (_,c):
            edit(c,self.note,event,1,{'title':'Edited','markdown':'Changed interpretation'})
            self.assertEqual(read(c,self.note)['researchProvenance'],self.payload['researchProvenance'])
            self.assertEqual(read(c,self.note,1)['researchProvenance'],self.payload['researchProvenance'])
            self.assertTrue(edit(c,self.note,event,1,{'title':'Edited','markdown':'Changed interpretation'})['replayed'])
        bad=copy.deepcopy(self.payload);bad['researchProvenance']['purpose']={'kind':'thesis_review'}
        with self.assertRaisesRegex(ValueError,'provenance'):self.save(2,str(uuid.uuid4()),bad)
        bad.pop('researchProvenance')
        with self.assertRaisesRegex(ValueError,'provenance'):self.save(2,str(uuid.uuid4()),bad)
