"""Safe export tests: extracted handlers, fake database, no email or iCloud writes."""
import ast
import base64
import copy
import io
import sys
import unittest
import uuid
from contextlib import contextmanager
from pathlib import Path
from flask import Flask, request, jsonify
from docx import Document
sys.path.insert(0,str(Path(__file__).resolve().parents[2]))
import summary_comparison

ROOT=Path(__file__).resolve().parents[2]
FIXTURE={'id':'version-one','summary_id':'note-one','status':'complete','baseline':{'title':'Fixture original'},'state':{'sections':{'brief':'Improved brief','takeaways':'Improved <script>literal</script> & 12%','questions':'Improved follow-ups','assessment':'Improved assessment'},'parts':{'1':{'record':'Second management passage'},'0':{'record':'First management passage'}}}}

class ExportTests(unittest.TestCase):
 def setUp(self):
  self.row=copy.deepcopy(FIXTURE);self.queued=[]
  owner=self
  class Cursor:
   def execute(self,sql,args):
    self.args=args;self.sql=sql
    if 'INSERT INTO icloud_export_tasks' in sql:owner.queued.append(args)
   def fetchone(self):
    return copy.deepcopy(owner.row) if self.args==('version-one','note-one') else None
  @contextmanager
  def db(commit=False):yield None,Cursor()
  self.app=Flask(__name__)
  env={'app':self.app,'request':request,'jsonify':jsonify,'get_db':db,'summary_comparison':summary_comparison,'io':io,'base64':base64,'uuid':uuid,'_safe_filename':lambda s,n:s[:n], '_VALID_SECTION_KEYS':{'all','brief','takeaways','meeting','questions','assessment','transcript'},'_SECTION_LABEL':{'all':'All','brief':'Brief'}}
  tree=ast.parse((ROOT/'app_v3.py').read_text())
  names={'_html_to_docx_elements','_generate_summary_docx_bytes','summary_section_to_docx','summary_save_to_icloud'}
  selected=[n for n in tree.body if isinstance(n,ast.FunctionDef) and n.name in names]
  exec(compile(ast.Module(body=selected,type_ignores=[]),'<export-handlers>','exec'),env)
  self.client=self.app.test_client()
 def test_download_uses_improved_only_and_literal_text(self):
  r=self.client.post('/api/summary-section-to-docx',json={'summaryId':'note-one','comparisonId':'version-one','section':'all'})
  self.assertEqual(r.status_code,200)
  self.assertIn('Improved',r.json['filename']);self.assertIn('version-one',r.json['filename'])
  doc=Document(io.BytesIO(base64.b64decode(r.json['fileData'])))
  text='\n'.join(p.text for p in doc.paragraphs)
  self.assertIn('Improved <script>literal</script> & 12%',text)
  self.assertLess(text.index('First management'),text.index('Second management'))
  self.assertIn('Improved follow-ups',text)
 def test_scoped_version_and_incomplete_rejected(self):
  body={'summaryId':'wrong-note','comparisonId':'version-one'}
  self.assertEqual(self.client.post('/api/summary-section-to-docx',json=body).status_code,404)
  self.row['status']='running';body['summaryId']='note-one'
  self.assertEqual(self.client.post('/api/summary-section-to-docx',json=body).status_code,409)
  self.assertEqual(self.client.post('/api/summaries/note-one/save-to-icloud',json=body).status_code,409)
  self.assertEqual(self.queued,[])
 def test_icloud_queues_exact_version_single_section(self):
  r=self.client.post('/api/summaries/note-one/save-to-icloud',json={'comparisonId':'version-one','section':'brief'})
  self.assertEqual(r.status_code,200);self.assertEqual(len(self.queued),1)
  doc=Document(io.BytesIO(base64.b64decode(self.queued[0][4])))
  text='\n'.join(p.text for p in doc.paragraphs)
  self.assertIn('Improved brief',text);self.assertNotIn('Improved follow-ups',text)
  self.assertIn('version-one',r.json['filename'])

if __name__=='__main__':unittest.main()
