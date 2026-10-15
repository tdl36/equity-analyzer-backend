import base64
import io
import unittest
import uuid
from flask import Flask,request
from tests.unit import test_pipeline_lab_note_jobs as fixtures
from tests.unit import test_stock_analysis_postgres as pg
from pipeline_lab_note_service import initialize
from pipeline_lab_note_review import create_blueprint,black_calibri
from pipeline_lab_notes import append

@unittest.skipUnless((pg.PG/'initdb').exists(),'Disposable PostgreSQL unavailable')
class Review(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        fixtures.Jobs.setUpClass.__func__(cls)
        with cls.db(True) as (_,c):initialize(c)
    tearDownClass=classmethod(fixtures.Jobs.tearDownClass.__func__)
    def setUp(self):
        fixtures.Jobs.setUp(self)
        with self.db(True) as (_,c):append(c,self.note,'ABT',str(uuid.uuid4()),0,self.payload)
        app=Flask('review')
        app.register_blueprint(create_blueprint(self.db,lambda:bool(request.headers.get('X-Auth')),
            lambda:'investor' if request.headers.get('X-Auth')=='investor' else None,lambda *args:''))
        self.client=app.test_client();self.path='/api/pipeline-lab/notes/'+self.note
    def test_acceptance_requires_investor_and_exact_revision_and_replays(self):
        for headers,expected in [({},401),({'X-Auth':'machine'},403)]:
            self.assertEqual(self.client.post(self.path+'/accept',headers=headers,json={'revision':1}).status_code,expected)
        headers={'X-Auth':'investor'}
        for _ in range(2):self.assertEqual(self.client.post(self.path+'/accept',headers=headers,json={'revision':1}).status_code,200)
        with self.db(True) as (_,c):append(c,self.note,'ABT',str(uuid.uuid4()),1,{**self.payload,'markdown':'Edited'})
        jobs=self.client.get('/api/pipeline-lab/note-jobs/ABT',headers=headers).json['jobs']
        self.assertEqual(jobs[0]['revision'],2);self.assertEqual(jobs[0]['acceptedRevision'],1)
        self.assertEqual(self.client.post(self.path+'/accept',headers=headers,json={'revision':3}).status_code,409)
    def test_exports_are_exact_revision_and_keep_caveats(self):
        headers={'X-Auth':'investor'}
        result=self.client.get(self.path+'/export?revision=1&format=md',headers=headers)
        self.assertEqual(result.status_code,200)
        self.assertIn(b'DRAFT',result.data);self.assertIn(b'SHA-256',result.data)
        self.assertIn('no-store',result.headers['Cache-Control'])
        self.assertEqual(self.client.get(self.path+'/export?revision=8&format=md',headers=headers).status_code,404)
        self.assertEqual(self.client.get(self.path+'/export?format=md',headers=headers).status_code,400)
    def test_all_word_runs_and_styles_black_calibri(self):
        from docx import Document
        from docx.shared import RGBColor
        doc=Document();doc.add_heading('Title');doc.add_table(rows=1,cols=1).cell(0,0).text='Table'
        doc.sections[0].header.paragraphs[0].add_run('Header').font.color.rgb=RGBColor(100,100,100)
        buf=io.BytesIO();doc.save(buf)
        result=Document(black_calibri(base64.b64encode(buf.getvalue())))
        for p in [*result.paragraphs,*result.sections[0].header.paragraphs,result.tables[0].cell(0,0).paragraphs[0]]:
            for run in p.runs:self.assertEqual(run.font.name,'Calibri');self.assertEqual(str(run.font.color.rgb),'000000')
