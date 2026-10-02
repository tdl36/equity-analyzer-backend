"""Hermetic Meeting Notes storage, generation, export and recovery checks."""
import ast
import io
import unittest
from pathlib import Path
from unittest.mock import Mock
import meeting_notes
import summary_lab
from summary_bulk import sections_html
from tests.unit.test_summary_bulk import ns

TREE = ast.parse(Path('app_v3.py').read_text())

class MeetingNotesTests(unittest.TestCase):
    def test_original_uses_source_and_rejects_empty_generation(self):
        scope={'meeting_notes':meeting_notes, '_call_llm_stream_with_retry':Mock(return_value={'text':'<ul><li>~5%, not committed</li></ul>'})}
        node=next(n for n in TREE.body if isinstance(n,ast.FunctionDef) and n.name=='_meeting_notes_html')
        exec(compile(ast.Module(body=[node],type_ignores=[]),'fixture','exec'),scope)
        self.assertIn('~5%',scope['_meeting_notes_html']('SOURCE TAIL: might fall', 'fake', korean=True))
        args=scope['_call_llm_stream_with_retry'].call_args.kwargs
        self.assertIn('SOURCE TAIL: might fall',args['messages'][0]['content'])
        self.assertIn('Write in Korean',args['messages'][0]['content'])
        scope['_call_llm_stream_with_retry'].return_value={'text':''}
        with self.assertRaises(ValueError):scope['_meeting_notes_html']('x')

    def test_notes_export_isolated_and_legacy_rows_compatible(self):
        from docx import Document
        row=summary_lab.export_row({'state':{'sections':{'notes':'- ~5%, not committed [P1]'}},'title':'Fixture'},'notes')
        self.assertIn('~5%',row['meeting_notes'])
        doc=Document(io.BytesIO(ns['_generate_summary_docx_bytes'](row,['notes'])))
        self.assertIn('not committed','\n'.join(p.text for p in doc.paragraphs))
        self.assertIn('Meeting Notes',sections_html(row,['notes']))
        self.assertNotIn('Meeting Notes',sections_html({'summary':'old'}))
        for p in doc.paragraphs:
            for r in p.runs:
                self.assertEqual(r.font.name,'Calibri')
                self.assertEqual(str(r.font.color.rgb),'000000')

    def test_all_insert_shapes_match(self):
        # Check actual query/tuple shapes without opening the application/database.
        checked=0
        for node in ast.walk(TREE):
            if not isinstance(node,ast.Call) or not isinstance(node.func,ast.Attribute) or node.func.attr!='execute' or len(node.args)<2: continue
            sql,params=node.args[:2]
            if isinstance(sql,ast.Constant) and isinstance(sql.value,str) and 'INSERT INTO meeting_summaries' in sql.value and 'meeting_notes' in sql.value:
                self.assertIsInstance(params,ast.Tuple)
                self.assertEqual(sql.value.count('%s'),len(params.elts),sql.value)
                checked+=1
        self.assertEqual(checked,4)

    def test_note_draft_is_checked_and_resumes_without_duplicate_calls(self):
        state={}; calls=[]
        def ask(system,prompt,tokens):
            calls.append(prompt)
            if 'Return JSON only:' in prompt:
                return '{"record":"might fall","passages":["might fall"],"issues":[]}'
            return '- might fall [P1]'
        summary_lab.generate('might fall',state,ask,lambda s:None)
        self.assertIn('notes',state['completedSections'])
        self.assertTrue(any('Review the draft against ORIGINAL' in p and '- might fall' in p for p in calls))
        count=len(calls)
        summary_lab.generate('might fall',state,ask,lambda s:None)
        self.assertEqual(count,len(calls))

class SavedNotesTests(unittest.TestCase):
    def run_request(self, rows, response_text='<ul><li>Saved</li></ul>'):
        from flask import Flask, request, jsonify
        from contextlib import contextmanager
        cur=Mock(); cur.fetchone.side_effect=rows
        @contextmanager
        def db(**kw): yield Mock(),cur
        fn=next(n for n in TREE.body if isinstance(n,ast.FunctionDef) and n.name=='generate_saved_meeting_notes')
        fn.decorator_list=[]
        scope=dict(request=request,jsonify=jsonify,get_db=db,cache=Mock(),_meeting_notes_html=Mock(return_value=response_text))
        exec(compile(ast.Module(body=[fn],type_ignores=[]),'fixture','exec'),scope)
        app=Flask(__name__)
        with app.test_request_context(json={}):
            response=scope['generate_saved_meeting_notes']('fixture')
        return response,scope,cur
    def test_duplicate_returns_saved_without_model(self):
        response,scope,cur=self.run_request([{'ok':True},{'meeting_notes':'existing'}])
        self.assertEqual(response.json['meetingNotes'],'existing')
        scope['_meeting_notes_html'].assert_not_called()
    def test_lock_conflict_never_calls_model(self):
        response,scope,cur=self.run_request([{'ok':False}])
        self.assertEqual(response[1],409)
        scope['_meeting_notes_html'].assert_not_called()
    def test_missing_source_never_calls_model(self):
        response,scope,cur=self.run_request([{'ok':True},{'raw_notes':''}])
        self.assertEqual(response[1],400)
        scope['_meeting_notes_html'].assert_not_called()
    def test_new_notes_update_only_new_section(self):
        response,scope,cur=self.run_request([{'ok':True},{'raw_notes':'original tail','source_meta':{}}])
        self.assertIn('Saved',response.json['meetingNotes'])
        self.assertEqual(cur.execute.call_args.args[0],'UPDATE meeting_summaries SET meeting_notes=%s WHERE id=%s')
