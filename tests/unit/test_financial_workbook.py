import copy
import unittest
from io import BytesIO
from openpyxl import load_workbook
from financial_workbook import export

class WorkbookTests(unittest.TestCase):
    def fixture(self):
        item=lambda v:{'value':v,'unit':'USD','tag':'RevenueFromContract','accession':'0000000001-26-000010','filed':'2026-02-01'}
        return {'version':3,'asOf':'2026-10-10','retrievedAt':'2026-10-10T12:00:00Z','cik':1,'annuals':[{'start':'2025-01-01','end':'2025-12-31','revenue':item(1000),'netIncome':item(0),'freeCashFlow':{**item(-25),'calculation':'operatingCashFlow - capex'},'operatingMargin':{**item(4.2),'unit':'%'}}],'financialWarnings':['=FORMULA_SHOULD_REMAIN_TEXT()']}
    def test_saved_values_zero_negatives_blanks_and_citations(self):
        snapshot=self.fixture();before=copy.deepcopy(snapshot);raw=export('SYNTH',snapshot,'assignment');book=load_workbook(BytesIO(raw));h=book['Financial history'];s=book['Source observations']
        self.assertEqual(book.sheetnames,['Financial history','Source observations']);self.assertEqual(h['C6'].value,1000);self.assertEqual(h['C11'].value,0);self.assertEqual(h['C13'].value,-25);self.assertEqual(h['C12'].value,4.2);self.assertIsNone(h['C7'].value)
        self.assertEqual(s['H7'].hyperlink.target,'https://www.sec.gov/Archives/edgar/data/1/000000000126000010/0000000001-26-000010-index.html')
        self.assertEqual(s['H8'].value,'Unavailable or withheld');self.assertEqual(snapshot,before)
        self.assertEqual(h.freeze_panes,'C6');self.assertEqual(s.freeze_panes,'D7');self.assertEqual(s.auto_filter.ref,'A6:H14')
        for sheet in book:
            for row in sheet:
                for cell in row:
                    if cell.value is not None:self.assertEqual(cell.font.name,'Calibri');self.assertEqual(cell.font.color.rgb,'00000000');self.assertNotEqual(cell.data_type,'f')
        warning=next(c for row in h for c in row if c.value=='=FORMULA_SHOULD_REMAIN_TEXT()');self.assertEqual(warning.data_type,'s')
    def test_nonfinite_boolean_unknown_and_oversized_are_not_numeric(self):
        f=self.fixture();f['annuals'][0]['revenue']['value']=float('nan');f['annuals'][0]['netIncome']['value']=True
        b=load_workbook(BytesIO(export('SYNTH',f,'assignment')));self.assertIsNone(b['Financial history']['C6'].value);self.assertIsNone(b['Financial history']['C11'].value)
        for value in [None,{}, {'annuals':[]},{'annuals':[{}]*6}]:
            with self.assertRaises(ValueError):export('SYNTH',value,'assignment')

    def test_incompatible_units_are_not_mislabeled(self):
        f=self.fixture();f['annuals'][0]['revenue']['unit']='EUR'
        b=load_workbook(BytesIO(export('SYNTH',f,'assignment')))
        self.assertIsNone(b['Financial history']['C6'].value)
        self.assertEqual(b['Source observations']['D7'].value,1000)
        self.assertEqual(b['Source observations']['E7'].value,'EUR')
