"""Authenticated Lab note list, revision-specific acceptance and portable export."""
import base64
import io
import re
from flask import Blueprint, jsonify, request, send_file
from pipeline_lab_notes import read, _note_identity


def accept(cur,note_id,revision,investor):
    note_id=_note_identity(note_id)
    if not investor:raise ValueError('Investor sign-in is required to accept a note.')
    if type(revision) is not int or revision<1:raise ValueError('An exact revision is required.')
    cur.execute('SELECT pg_advisory_xact_lock(hashtext(%s))',('lab-note:'+note_id,))
    cur.execute('SELECT accepted_at FROM pipeline_lab_note_acceptances WHERE note_id=%s AND revision=%s',(note_id,revision))
    prior=cur.fetchone()
    if prior:return {'noteId':note_id,'revision':revision,'acceptedAt':prior['accepted_at'].isoformat()}
    draft=read(cur,note_id)
    if not draft or draft['revision']!=revision:raise ValueError('The draft changed. Review the current revision before acceptance.')
    cur.execute('''INSERT INTO pipeline_lab_note_acceptances(note_id,revision,investor)
        VALUES(%s,%s,%s) RETURNING accepted_at''',(note_id,revision,investor))
    return {'noteId':note_id,'revision':revision,'acceptedAt':cur.fetchone()['accepted_at'].isoformat()}


def export_text(draft,accepted=False):
    evidence=draft.get('generationEvidence') or {}
    parts=[draft['markdown'],'\n## Review record',
           ('Investor accepted' if accepted else 'DRAFT — not investor accepted')+f" · revision {draft['revision']} · saved {draft['createdAt']}",
           '\n## Quality and source limitations',*['- '+s for s in draft['qualityIssues']],
           '\n## Sources appendix',evidence.get('sourcesMarkdown') or 'No sources appendix supplied.']
    from pipeline_lab_note_provenance import export_lines
    parts.extend(export_lines(draft.get('researchProvenance')))
    parts.append('\n## Frozen original manifest')
    for s in draft['sourcePack']:parts.append(f"- {s['filename']} · {s['reference']} · {s['bytes']} bytes · SHA-256 {s['sha256']}")
    for chart in evidence.get('chartSpecs',[]):
        parts.extend(['\n### '+chart['period']+' — '+chart['kind'],'Units and figures require source verification.','| Segment | Value |','| --- | ---: |'])
        parts.extend('| '+r['segment'].replace('|','/')+' | '+str(r['value'])+' |' for r in chart['data'])
    return '\n'.join(parts)


def black_calibri(encoded):
    from docx import Document
    from docx.shared import RGBColor
    from docx.oxml.ns import qn
    doc=Document(io.BytesIO(base64.b64decode(encoded)))
    for style in doc.styles:
        if hasattr(style,'font'):style.font.name='Calibri';style.font.color.rgb=RGBColor(0,0,0)
    # Every run, including table/header/footer text and existing hyperlink runs.
    for part in [doc.part,*[p for p in doc.part.related_parts.values() if hasattr(p,'element')]]:
        root=getattr(part,'element',None)
        if root is None:continue
        for run in root.iter(qn('w:r')):
            from docx.text.run import Run
            font=Run(run,None).font;font.name='Calibri';font.color.rgb=RGBColor(0,0,0)
    out=io.BytesIO();doc.save(out);out.seek(0);return out


def word_export(ticker,title,text):
    from bs4 import BeautifulSoup, NavigableString
    from markdown import markdown
    from docx import Document
    from docx.shared import Pt, Inches
    doc=Document()
    doc.sections[0].top_margin=Inches(.7);doc.sections[0].bottom_margin=Inches(.7)
    doc.styles['Normal'].font.size=Pt(11)
    doc.add_heading(title,0)
    soup=BeautifulSoup(markdown(text,extensions=['tables']),'html.parser')
    def runs(paragraph,node,bold=False,italic=False):
        if isinstance(node,NavigableString):
            r=paragraph.add_run(str(node));r.bold=bold;r.italic=italic;return
        for child in node.children:runs(paragraph,child,bold or node.name in ('strong','b'),italic or node.name in ('em','i'))
    for node in soup.children:
        name=getattr(node,'name',None)
        if not name:continue
        if name=='table':
            rows=node.find_all('tr')
            count=max((len(r.find_all(['td','th'])) for r in rows),default=0)
            if not count:continue
            table=doc.add_table(rows=0,cols=count);table.style='Light Shading Accent 1'
            for row in rows:
                cells=table.add_row().cells
                for cell,item in zip(cells,row.find_all(['td','th'])):runs(cell.paragraphs[0],item,item.name=='th')
        elif name in ('ul','ol'):
            for item in node.find_all('li',recursive=False):runs(doc.add_paragraph(style='List Bullet' if name=='ul' else 'List Number'),item)
        elif name in ('h1','h2','h3','h4','h5','h6'):runs(doc.add_heading(level=min(int(name[1]),3)),node)
        else:runs(doc.add_paragraph(),node)
    buf=io.BytesIO();doc.save(buf)
    return black_calibri(base64.b64encode(buf.getvalue()))


def create_blueprint(get_db,authenticated,investor_identity,docx_export):
    bp=Blueprint('pipeline_lab_note_review',__name__)
    @bp.before_request
    def auth():
        if not authenticated():return jsonify(error='Sign in to review Lab notes.'),401
    @bp.after_request
    def private(response):
        response.headers['Cache-Control']='private, no-store'
        response.headers['X-Content-Type-Options']='nosniff'
        return response
    @bp.get('/api/pipeline-lab/note-jobs/<ticker>')
    def listing(ticker):
        if not re.fullmatch('[A-Z0-9][A-Z0-9.-]{0,19}',ticker):return jsonify(error='Invalid ticker.'),400
        with get_db() as (_,cur):
            cur.execute('''SELECT j.note_id,j.status,j.reason,j.created_at,
                (SELECT MAX(revision) FROM pipeline_lab_note_versions v WHERE v.note_id=j.note_id) AS revision,
                (SELECT MAX(revision) FROM pipeline_lab_note_acceptances a WHERE a.note_id=j.note_id) AS accepted_revision
                FROM pipeline_lab_note_jobs j JOIN pipeline_lab_note_plans p USING(note_id)
                WHERE p.plan->>'ticker'=%s ORDER BY j.created_at DESC LIMIT 30''',(ticker,))
            rows=[{'noteId':str(r['note_id']),'status':r['status'],'reason':r['reason'],
                   'revision':r['revision'],'acceptedRevision':r['accepted_revision'],
                   'createdAt':r['created_at'].isoformat()} for r in cur.fetchall()]
        return jsonify(jobs=rows)
    @bp.post('/api/pipeline-lab/notes/<note_id>/accept')
    def acceptance(note_id):
        if not investor_identity():return jsonify(error='Investor sign-in is required; machine credentials cannot accept research.'),403
        if request.content_length is None or request.content_length>512:return jsonify(error='Bounded revision required.'),400
        data=request.get_json(silent=True)
        if not isinstance(data,dict) or set(data)!={'revision'}:return jsonify(error='Exact revision required.'),400
        try:
            with get_db(commit=True) as (_,cur):result=accept(cur,note_id,data['revision'],investor_identity())
            return jsonify(result)
        except ValueError as exc:return jsonify(error=str(exc)),409
    @bp.get('/api/pipeline-lab/notes/<note_id>/export')
    def export(note_id):
        try:
            revision=request.args.get('revision','');fmt=request.args.get('format','docx')
            if not re.fullmatch('[1-9][0-9]{0,8}',revision) or fmt not in ('docx','md'):raise ValueError('Choose an exact revision and Word or Markdown format.')
            with get_db() as (_,cur):
                draft=read(cur,note_id,int(revision))
                if not draft:return jsonify(error='Draft revision not found.'),404
                cur.execute('SELECT 1 FROM pipeline_lab_note_acceptances WHERE note_id=%s AND revision=%s',(note_id,int(revision)))
                accepted=bool(cur.fetchone())
            text=export_text(draft,accepted)
            if fmt=='md':data=io.BytesIO(text.encode());mime='text/markdown; charset=utf-8'
            else:
                data=word_export(draft['ticker'],draft['title'],text);mime='application/vnd.openxmlformats-officedocument.wordprocessingml.document'
            return send_file(data,mimetype=mime,as_attachment=True,download_name=f"{draft['ticker']}-Lab-note-v{revision}.{fmt}")
        except ValueError as exc:return jsonify(error=str(exc)),400
    return bp
