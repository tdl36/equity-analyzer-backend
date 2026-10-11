"""Production export of frozen financial observations; no model or network calls."""
from io import BytesIO
from datetime import date
import math
import re
from openpyxl import Workbook
from openpyxl.styles import Alignment,Font,PatternFill,Border,Side
from openpyxl.utils import get_column_letter
from financial_provenance import LABELS


def number(value):return isinstance(value,(int,float)) and not isinstance(value,bool) and math.isfinite(value)


def compatible(key,item):
    expected="USD/shares" if key=="eps" else "%" if key=="operatingMargin" else "USD"
    return isinstance(item,dict) and item.get("unit")==expected and number(item.get("value"))


def export(ticker,snapshot,assignment_id):
    if not isinstance(snapshot,dict) or not snapshot.get('annuals'):raise ValueError('No saved annual financial observations are available for this assignment.')
    rows=snapshot['annuals']
    if not isinstance(rows,list) or len(rows)>5 or any(not isinstance(r,dict) for r in rows):raise ValueError('The saved financial snapshot has an unsupported format.')
    # Export the retained observations, not refreshed prices or a new forecast model.
    book=Workbook();history=book.active;history.title='Financial history';sources=book.create_sheet('Source observations')
    font=Font(name='Calibri',size=11,color='000000');bold=Font(name='Calibri',size=11,bold=True,color='000000')
    fill=PatternFill('solid',fgColor='E8F1F5');line=Side(style='thin',color='D8E1E5')
    def put(ws,r,c,value,bolded=False):
        if isinstance(value,str):value=re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f]','',value)
        cell=ws.cell(r,c,value)
        if isinstance(value,str):cell.data_type='s'  # A source string must never become an Excel formula.
        cell.font=bold if bolded else font;cell.alignment=Alignment(vertical='top',wrap_text=True)
        return cell
    for ws in (history,sources):
        ws.sheet_view.showGridLines=False
        put(ws,1,1,ticker+' financial history',True);ws['A1'].font=Font(name='Calibri',size=20,bold=True,color='000000');ws.row_dimensions[1].height=30
        ws.merge_cells(start_row=1,start_column=1,end_row=1,end_column=8)
        put(ws,2,1,'Saved observations as of '+str(snapshot.get('asOf','unavailable'))+'. Exporting does not refresh the source data.');ws.merge_cells('A2:H2');ws.row_dimensions[2].height=32
        put(ws,3,1,'Blank values are unavailable, withheld or have incompatible units. This workbook contains saved values, not a live financial or forecast model.');ws.merge_cells('A3:H3');ws.row_dimensions[3].height=32
        ws.sheet_properties.pageSetUpPr.fitToPage=True;ws.page_setup.orientation='landscape';ws.page_setup.paperSize=ws.PAPERSIZE_A3;ws.page_setup.fitToWidth=1;ws.page_setup.fitToHeight=0
        ws.oddFooter.center.text='Page &P of &N';ws.oddFooter.center.font='Calibri';ws.oddFooter.center.size=9;ws.oddFooter.center.color='000000'
    put(history,5,1,'Metric',True);put(history,5,2,'Units',True)
    for col,row in enumerate(rows,3):
        value=row.get('end') or row.get('year') or 'Period unavailable'
        try:value=date.fromisoformat(value)
        except (ValueError,TypeError):pass
        c=put(history,5,col,value,True);c.number_format='yyyy-mm-dd'
        history.column_dimensions[get_column_letter(col)].width=19
    for r,(key,label) in enumerate(LABELS.items(),6):
        put(history,r,1,label,True);unit='USD per share' if key=='eps' else '%' if key=='operatingMargin' else 'USD'
        put(history,r,2,unit)
        for col,row in enumerate(rows,3):
            item=row.get(key);value=item.get('value') if isinstance(item,dict) else None
            c=put(history,r,col,value if compatible(key,item) else None)
            c.number_format='#,##0.00;(#,##0.00);0.00' if key=='eps' else '0.0"%";(0.0"%");0.0"%"' if key=='operatingMargin' else '#,##0;(#,##0);0'
            c.alignment=Alignment(horizontal='right',vertical='top')
    put(history,17,1,'Definitions and limitations',True);history.merge_cells('A17:H17')
    notes=['Amounts are unscaled source units. Operating margin values are percentage points (4.2 means 4.2%).',
           'Consolidated free cash flow = operating cash flow less reported capital expenditure. It may differ from industrial or distributable cash flow.',
           snapshot.get('historyBasis') or 'Saved annual observations may include later restatements. Check filing dates and definitions on Source observations.',
           'Source observations contains the exact values, filing dates, definitions and filing links used for this export.']
    notes+=list(snapshot.get('financialWarnings') or [])+list(snapshot.get('gaps') or [])
    for i,note in enumerate(notes,18):put(history,i,1,str(note));history.merge_cells(start_row=i,start_column=1,end_row=i,end_column=8);history.row_dimensions[i].height=34
    history.column_dimensions['A'].width=27;history.column_dimensions['B'].width=18;history.freeze_panes='C6';history.print_title_rows='1:5';history.print_options.horizontalCentered=True
    meta='Retrieved '+str(snapshot.get('retrievedAt','unavailable'))+' · Snapshot version '+str(snapshot.get('version','unavailable'))+' · Assignment '+str(assignment_id)
    put(sources,4,1,meta);sources.merge_cells('A4:H4');sources.row_dimensions[4].height=32
    headers=['Period start','Period end','Metric','Value (unscaled)','Unit','Source tag / calculation','Filed','Filing / availability']
    for col,label in enumerate(headers,1):put(sources,6,col,label,True)
    r=7
    for period in rows:
        for key,label in LABELS.items():
            item=period.get(key) if isinstance(period.get(key),dict) else {};value=item.get('value')
            fields=[period.get('start'),period.get('end'),label,value if number(value) else None,item.get('unit'),item.get('calculation') or item.get('tag'),item.get('filed'),'Unavailable or withheld' if not number(value) else item.get('accession') or 'Accession unavailable']
            for col,value in enumerate(fields,1):
                if col in (1,2,7) and value:
                    try:value=date.fromisoformat(value)
                    except (ValueError,TypeError):pass
                c=put(sources,r,col,value)
                if col in (1,2,7):c.number_format='yyyy-mm-dd'
                if col==4:c.number_format='#,##0.00;(#,##0.00);0.00'
            acc=str(item.get('accession',''));cik=str(snapshot.get('cik',''))
            if number(item.get('value')) and re.fullmatch(r'\d{10}-\d{2}-\d{6}',acc) and re.fullmatch(r'\d{1,10}',cik):
                sources.cell(r,8).hyperlink=f'https://www.sec.gov/Archives/edgar/data/{int(cik)}/{acc.replace("-", "")}/{acc}-index.html'
                sources.cell(r,8).font=Font(name='Calibri',size=11,color='000000',underline='single')
            sources.row_dimensions[r].height=42;r+=1
    sources.auto_filter.ref=f'A6:H{r-1}';sources.freeze_panes='D7';sources.print_title_rows='1:6'
    for col,width in zip('ABCDEFGH',[14,14,25,23,16,44,14,34]):sources.column_dimensions[col].width=width
    for ws,header in [(history,5),(sources,6)]:
        for cell in ws[header]:cell.fill=fill;cell.border=Border(bottom=line)
        ws.print_area=f'A1:{get_column_letter(max(8,ws.max_column))}{ws.max_row}'
    book.calculation.fullCalcOnLoad=False  # No formulas, external links or macros are authored.
    out=BytesIO();book.save(out);return out.getvalue()
