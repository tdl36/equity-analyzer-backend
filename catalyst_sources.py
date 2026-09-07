"""Shared local catalyst source inventory; never includes generated recaps."""
import hashlib
import json
from pathlib import Path

SOURCE_EXTS = {'.pdf', '.docx', '.xlsx', '.xlsm', '.txt', '.md', '.csv', '.tsv', '.pptx', '.html', '.htm'}
GENERATED_PREFIXES = ('recap_', 'synthesis_')


def inventory(folder, excluded=()):
    root = Path(folder).resolve()
    if not root.is_dir():
        raise ValueError(f'Event folder is missing: {root.name}')
    excluded = set(excluded)
    files, issues = [], []
    for path in sorted(root.rglob('*')):
        rel = path.relative_to(root)
        if any(p.startswith('.') or p.startswith('~$') or p.lower() == 'processed' for p in rel.parts):
            # Cloud-only placeholders must be visible as a blocking issue.
            if path.name.endswith('.icloud'):
                issues.append(f'{rel}: iCloud download pending; choose Download Now in Finder')
            continue
        if path.is_symlink() or not path.resolve().is_relative_to(root):
            continue
        if not path.is_file() or path.name.lower().startswith(GENERATED_PREFIXES):
            continue
        if rel.as_posix() in excluded or path.name in excluded:
            continue
        if path.suffix.lower() not in SOURCE_EXTS:
            if path.suffix.lower() in {'.doc', '.xls', '.ppt', '.zip'}:
                issues.append(f'{rel}: convert or extract this file before generating the recap')
            continue
        stat = path.stat()
        if not stat.st_size:
            issues.append(f'{rel}: empty file or incomplete iCloud download')
        else:
            files.append({'path': path, 'name': rel.as_posix(), 'size': stat.st_size, 'mtime_ns': stat.st_mtime_ns})
    return files, issues


def fingerprint(files):
    return hashlib.sha256(json.dumps([(f['name'], f['size'], f['mtime_ns']) for f in files], separators=(',', ':')).encode()).hexdigest()


def read_sources(folder, excluded=()):
    """Read all selected sources, or fail before paid generation on partial input."""
    import base64
    import io
    import zipfile
    from xml.etree import ElementTree as ET
    files, issues = inventory(folder, excluded)
    parts = []
    for item in files:
        path, name = item['path'], item['name']
        try:
            data = path.read_bytes()
            ext = path.suffix.lower()
            if ext == '.pdf':
                from PyPDF2 import PdfReader
                pdf = PdfReader(io.BytesIO(data))
                if pdf.is_encrypted or not len(pdf.pages):
                    raise ValueError('encrypted PDF or no pages')
                parts.append({'type': 'pdf', 'name': name, 'data': base64.b64encode(data).decode('ascii')})
                continue
            if ext == '.docx':
                from docx import Document
                doc = Document(io.BytesIO(data))
                text = '\n'.join([p.text for p in doc.paragraphs] + ['\t'.join(c.text for c in row.cells) for table in doc.tables for row in table.rows])
            elif ext in {'.xlsx', '.xlsm'}:
                import openpyxl
                wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True, read_only=True)
                try:
                    text = '\n'.join(f'Sheet: {ws.title}\n' + '\n'.join('\t'.join('' if c is None else str(c) for c in row) for row in ws.iter_rows(values_only=True)) for ws in wb.worksheets)
                finally:
                    wb.close()
            elif ext == '.pptx':
                with zipfile.ZipFile(io.BytesIO(data)) as archive:
                    slides = sorted((n for n in archive.namelist() if __import__('re').fullmatch(r'ppt/slides/slide\d+\.xml', n)), key=lambda n: int(__import__('re').search(r'slide(\d+)', n).group(1)))
                    text = '\n'.join(f'Slide {i+1}\n' + '\n'.join(e.text or '' for e in ET.fromstring(archive.read(n)).iter() if e.tag.endswith('}t')) for i, n in enumerate(slides))
            elif ext in {'.html', '.htm'}:
                from bs4 import BeautifulSoup
                soup = BeautifulSoup(data, 'html.parser')
                for e in soup(['script', 'style']):
                    e.decompose()
                text = soup.get_text('\n', strip=True)
            else:
                text = data.decode('utf-8-sig')
            if not text.strip():
                raise ValueError('no extractable text; export to PDF or OCR this document')
            parts.append({'type': 'text', 'name': name, 'content': text})
        except Exception as exc:
            issues.append(f'{name}: {type(exc).__name__}: {exc}')
    if issues:
        raise ValueError('Source preflight stopped; no model call made. ' + '; '.join(issues[:10]))
    if not parts:
        raise ValueError('No supported source documents in this event folder or its subfolders. Add PDF, DOCX, PPTX, XLSX, text or HTML sources; generated recaps are excluded.')
    return parts
