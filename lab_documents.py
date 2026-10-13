"""Bounded, loss-aware document intake. Preparation never calls a model."""
import base64
import hashlib
import io
from pathlib import Path
import zipfile
import xml.etree.ElementTree as ET
from PyPDF2 import PdfReader, PdfWriter
from PIL import Image

MAX_BYTES = 20 * 1024 * 1024
MAX_PARTS = 40
SUPPORTED = '.pdf, .png, .jpg, .jpeg, .webp, .gif, .docx, .pptx, .xlsx, .txt, .md, .csv, .tsv, .json, .html'

def block(kind, mime, data):
    return {'type':kind,'source':{'type':'base64','media_type':mime,'data':base64.b64encode(data).decode()}}

def text_parts(text, label):
    from summary_comparison import split_text
    if not text.strip(): return []
    return [{'label':label+f' · text {a+1}–{b}', 'text':body,'blocks':[], 'visual':False} for a,b,body in split_text(text,20000)]

def image_part(data, label):
    if len(data)>5*1024*1024: raise ValueError('An image exceeds 5 MB. Supply a smaller original or PDF.')
    with Image.open(io.BytesIO(data)) as im:
        if getattr(im,'n_frames',1)>1: raise ValueError('Animated images are not supported; choose individual frames.')
        if max(im.size)>8000: raise ValueError('An image exceeds 8,000 pixels; supply a smaller image or PDF.')
        mime=Image.MIME.get(im.format)
        if mime not in ('image/png','image/jpeg','image/webp','image/gif'): raise ValueError('Embedded image format needs conversion to PNG/JPEG or a PDF export.')
        im.verify()
    return {'label':label,'text':'','blocks':[block('image',mime,data)],'visual':True}

def prepare(filename, data):
    filename=Path(filename).name
    if not filename or not data or len(data)>MAX_BYTES: raise ValueError('Choose a nonempty file of at most 20 MB.')
    ext=Path(filename).suffix.lower(); parts=[]; limitations=[]; pages=None
    if ext=='.pdf':
        reader=PdfReader(io.BytesIO(data))
        if reader.is_encrypted: raise ValueError('Unlock this PDF before importing it.')
        pages=len(reader.pages)
        if not 0<pages<=200: raise ValueError('Choose a PDF with 1–200 pages; split longer documents explicitly.')
        for start in range(0,pages,5):
            writer=PdfWriter(); texts=[]
            for page in reader.pages[start:start+5]:
                writer.add_page(page)
                try:texts.append(page.extract_text() or '')
                except Exception:texts.append('')
            out=io.BytesIO();writer.write(out)
            if len(out.getvalue())>MAX_BYTES:raise ValueError('A five-page PDF part exceeds 20 MB. Split or optimize the PDF first.')
            parts.append({'label':f'{filename} · PDF pages {start+1}–{min(start+5,pages)}','text':'\n'.join(texts),'blocks':[block('document','application/pdf',out.getvalue())],'visual':True})
        limitations=['All PDF pages supplied as visual documents. Small print, complex charts and OCR may still require investor verification.']
    elif ext in ('.png','.jpg','.jpeg','.webp','.gif'):
        parts=[image_part(data,filename)];limitations=['The original image is reviewed visually; interpretation is model-generated.']
    elif ext in ('.docx','.pptx','.xlsx'):
        with zipfile.ZipFile(io.BytesIO(data)) as z:
            entries=z.infolist()
            if len(entries)>3000 or sum(e.file_size for e in entries)>80*1024*1024: raise ValueError('Office package exceeds the safe extraction bound.')
            names=z.namelist()
            if any('/embeddings/' in n or '/diagrams/' in n for n in names): raise ValueError('This Office file has embedded objects or diagrams. Export it to PDF for visual coverage.')
            if ext=='.xlsx':
                from openpyxl import load_workbook
                formulas=load_workbook(io.BytesIO(data),read_only=True,data_only=False)
                cached=load_workbook(io.BytesIO(data),read_only=True,data_only=True)
                try:
                    for sheet in formulas:
                        if sheet.max_row*sheet.max_column>200000:raise ValueError('Worksheet exceeds 200,000 cells. Supply a scoped workbook or PDF.')
                        lines=[]
                        for row, values in zip(sheet.iter_rows(),cached[sheet.title].iter_rows()):
                            for cell,value in zip(row,values):
                                if cell.value is not None:
                                    lines.append(f'{cell.coordinate}: {cell.value}'+(f' [cached value: {value.value}; not recalculated]' if cell.data_type=='f' else ''))
                        parts.extend(text_parts('\n'.join(lines),filename+' · worksheet '+sheet.title))
                finally:formulas.close();cached.close()
            # Include source XML text, table cells, formulas, chart caches and notes.
            # Internal locations remain visible; never imply rendered page fidelity.
            for name in sorted(names):
                if name.endswith('.xml') and ((ext=='.docx' and name.startswith('word/')) or (ext=='.pptx' and (name.startswith('ppt/slides/') or name.startswith('ppt/notesSlides/') or name.startswith('ppt/charts/'))) or (ext=='.xlsx' and name.startswith('xl/charts/'))):
                    root=ET.fromstring(z.read(name));values=[]
                    for e in root.iter():
                        tag=e.tag.rsplit('}',1)[-1]
                        if e.text and e.text.strip() and tag in ('t','v','f','sheet'):
                            values.append((('formula: ' if tag=='f' else '')+e.text).strip())
                    if values:parts.extend(text_parts('\n'.join(values),filename+' · '+name))
                if '/media/' in name and not name.endswith('/'):
                    parts.append(image_part(z.read(name),filename+' · '+name))
        limitations=['Office text, table values, formulas, chart data and embedded images are reviewed separately. Page/slide layout and chart appearance are not rendered; export to PDF for visual fidelity. Spreadsheet cells retain sheet names and coordinates; formula results are cached, not recalculated.']
    elif ext in ('.txt','.md','.csv','.tsv','.json','.html'):
        text=data.decode('utf-8-sig')
        if '\x00' in text:raise ValueError('This is not a supported UTF-8 text file.')
        parts=text_parts(text,filename)
        limitations=['Text/markup is reviewed completely; HTML linked images, scripts and external resources are not fetched.']
    else:raise ValueError('Unsupported document format. Supported: '+SUPPORTED+'. Use Audio for recordings; export other visual formats to PDF.')
    if not parts or len(parts)>MAX_PARTS:raise ValueError('Document must contain 1–40 review parts. Split larger files explicitly; nothing was truncated.')
    return {'filename':filename,'sha256':hashlib.sha256(data).hexdigest(),'bytes':len(data),'pages':pages,'parts':parts,'limitations':limitations}

def manifest(doc):
    return {k:v for k,v in doc.items() if k!='parts'} | {'processingVersion':1,'preparedHash':hashlib.sha256(__import__('json').dumps(doc['parts'],sort_keys=True).encode()).hexdigest(),'partCount':len(doc['parts']),'visualParts':sum(bool(p['visual']) for p in doc['parts'])}

def content(part, prompt):
    return [{'type':'text','text':part['label']},*part['blocks'],{'type':'text','text':prompt}] if part['visual'] else prompt

def validate_visual(review):
    if not isinstance(review,dict) or not isinstance(review.get('record'),str) or not review['record'].strip():raise ValueError('Visual review omitted its record.')
    if not isinstance(review.get('issues',[]),list):raise ValueError('Visual review issues must be a list.')
    return dict(review,passages=[],issues=list(review.get('issues',[]))+['Visual observations are model interpretations; quotations were not independently text-matched.'])
